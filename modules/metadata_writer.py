import os

from PIL import Image, PngImagePlugin
from pypdf import PdfReader, PdfWriter
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

def get_modified_path(file_path):
    base_name, extension = os.path.splitext(file_path)
    output_path = f"{base_name}_modified{extension}"
    counter = 1

    while os.path.exists(output_path):
        output_path = f"{base_name}_modified_{counter}{extension}"
        counter += 1

    return output_path


def write_jpeg_metadata(file_path, output_path, artist, description):
    with Image.open(file_path) as image:
        exif = image.getexif()
        exif[315] = artist
        exif[270] = description
        image.save(output_path, exif=exif)


def write_png_metadata(file_path, output_path, author, description):
    with Image.open(file_path) as image:
        png_info = PngImagePlugin.PngInfo()

        for key, value in image.info.items():
            if isinstance(value, str):
                png_info.add_text(key, value)

        png_info.add_text("Author", author)
        png_info.add_text("Description", description)
        image.save(output_path, pnginfo=png_info)


def write_pdf_metadata(file_path, output_path, title, author, subject, keywords):
    reader = PdfReader(file_path)
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    writer.add_metadata({
        "/Title": title,
        "/Author": author,
        "/Subject": subject,
        "/Keywords": keywords
    })

    with open(output_path, "wb") as file:
        writer.write(file)


def write_docx_metadata(file_path, output_path, title, author, subject, keywords):
    document = Document(file_path)
    properties = document.core_properties
    properties.title = title
    properties.author = author
    properties.subject = subject
    properties.keywords = keywords
    document.save(output_path)


def write_xlsx_metadata(file_path, output_path, title, author, subject, keywords):
    workbook = load_workbook(file_path)
    properties = workbook.properties
    properties.title = title
    properties.creator = author
    properties.subject = subject
    properties.keywords = keywords
    workbook.save(output_path)


def write_pptx_metadata(file_path, output_path, title, author, subject, keywords):
    presentation = Presentation(file_path)
    properties = presentation.core_properties
    properties.title = title
    properties.author = author
    properties.subject = subject
    properties.keywords = keywords
    presentation.save(output_path)


def metadata_writer():
    print("\n========================")
    print("Metadata Writer")
    print("========================")

    file_path = input("Enter file path: ").strip()

    if not os.path.isfile(file_path):
        print("File not found.")
        return

    extension = os.path.splitext(file_path)[1].lower()
    supported_formats = {
        ".jpg": "JPEG",
        ".jpeg": "JPEG",
        ".png": "PNG",
        ".pdf": "PDF",
        ".docx": "Word",
        ".xlsx": "Excel",
        ".pptx": "PowerPoint"
    }

    if extension not in supported_formats:
        print("Unsupported file type.")
        print("\nSupported formats:")
        print("JPEG, PNG, PDF, DOCX, XLSX, PPTX")
        return

    file_type = supported_formats[extension]
    print(f"Detected file type: {file_type}")
    print("\nEnter metadata")

    if file_type == "JPEG":
        artist = input("Artist: ").strip()
        description = input("Description: ").strip()
        writer = write_jpeg_metadata
        metadata = (artist, description)
    elif file_type == "PNG":
        author = input("Author: ").strip()
        description = input("Description: ").strip()
        writer = write_png_metadata
        metadata = (author, description)
    else:
        title = input("Title: ").strip()
        author = input("Author: ").strip()
        subject = input("Subject: ").strip()
        keywords = input("Keywords: ").strip()
        metadata = (title, author, subject, keywords)

        if file_type == "PDF":
            writer = write_pdf_metadata
        elif file_type == "Word":
            writer = write_docx_metadata
        elif file_type == "Excel":
            writer = write_xlsx_metadata
        else:
            writer = write_pptx_metadata

    output_path = get_modified_path(file_path)

    try:
        writer(file_path, output_path, *metadata)
        print("\nMetadata written successfully.")
        print(f"Original preserved: {file_path}")
        print(f"Modified copy:      {output_path}")
    except PermissionError as error:
        print(f"Could not write metadata: Permission denied: {error}")
    except Exception as error:
        print(f"Could not write metadata: {error}")
